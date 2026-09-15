import json
from pathlib import Path
import numpy as np
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parents[1]; D=ROOT/'docs'; m=json.loads((ROOT/'backend/models/fabric_mobilenetv2.manifest.json').read_text()); cm=m['confusion_matrix']; labels=[x.title() for x in cm['classes']]; arr=np.array(cm['matrix'])
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18);bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',20);small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',14);cell=86;left=135;top=90;img=Image.new('RGB',(left+cell*6+40,top+cell*6+70),'white');draw=ImageDraw.Draw(img);draw.text((left+135,20),'Confusion Matrix',font=bold,fill='#10231d');draw.text((left+125,53),'Predicted Fabric Class',font=small,fill='#38564a')
for j,label in enumerate(labels):draw.text((left+j*cell+5,top-25),label[:9],font=small,fill='#10231d')
draw.text((12,top-25),'True label',font=small,fill='#38564a')
maxv=arr.max()
for i,row in enumerate(arr):
 draw.text((left-95,top+i*cell+30),labels[i],font=small,fill='#10231d')
 for j,value in enumerate(row):
  shade=int(238-(value/maxv)*150);draw.rectangle((left+j*cell,top+i*cell,left+(j+1)*cell,top+(i+1)*cell),fill=(shade,245,shade),outline='#ffffff');draw.text((left+j*cell+34,top+i*cell+31),str(value),font=font,fill='#10231d')
png=D/'confusion_matrix_results.png';img.save(png)
doc=Document();sec=doc.sections[0];sec.left_margin=Inches(.8);sec.right_margin=Inches(.8);sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(11)
def para(text='',bold=False,center=False):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY;r=p.add_run(text);r.bold=bold;return p
def caption(text):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(text);r.italic=True;r.font.size=Pt(10)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('RESULTS AND DISCUSSION');r.bold=True;r.font.size=Pt(14)
para(f"The proposed LaundraAI system was evaluated on a held-out test set of {m['test_images']} images. The MobileNetV2 transfer-learning model obtained {m['test_accuracy']*100:.2f}% accuracy and {m['macro_f1']*100:.2f}% macro F1-score. These findings represent performance on the prepared dataset and should be interpreted as prototype evaluation results.")
doc.add_picture(str(png),width=Inches(5.9));caption('Fig. 5. Confusion matrix of the proposed fabric classification model on the held-out test set.')
para('Fig. 5 shows that the classifier correctly recognized most instances in every supported category. Denim produced the strongest class-wise F1-score (86.31%), indicating that its characteristic texture was comparatively easier for the model to distinguish. Silk produced the lowest F1-score (75.27%), indicating greater confusion with visually similar or reflective fabric surfaces. The non-fabric class also reduced the likelihood that unrelated images were treated as garment textiles.')
para('Table I summarizes the overall classification performance. The recorded model inference time was 0.8267 seconds in the evaluation artifact, which supports interactive use in a web-based planning prototype.')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid';heads=['Accuracy','Precision','Recall','F1-score','Response time'];vals=[f"{m['test_accuracy']*100:.2f}%",'80.85%',f"{m['test_accuracy']*100:.2f}%",f"{m['macro_f1']*100:.2f}%",f"{m['test_inference_seconds']:.4f} s"]
for c,h in zip(t.rows[0].cells,heads):c.text=h
for c,v in zip(t.add_row().cells,vals):c.text=v
caption('Table I. Overall fabric classification performance.')
para('Fabric-wise results show that all classes achieved an F1-score above 75%. The model performed well on cotton, polyester, denim and wool, while silk remains a priority category for additional dataset collection and model refinement.')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid';heads=['Fabric','Support','Precision','Recall','F1-score'];
for c,h in zip(t.rows[0].cells,heads):c.text=h
for name,v in m['per_class_metrics'].items():
 cells=t.add_row().cells
 for c,x in zip(cells,[name.title(),v['support'],f"{v['precision']*100:.2f}%",f"{v['recall']*100:.2f}%",f"{v['f1_score']*100:.2f}%"]):c.text=str(x)
caption('Table II. Fabric-wise classification results.')
para('Fig. 6 should present actual screenshots from the LaundraAI result page after uploading verified garment images. Each screenshot should show the uploaded image, predicted fabric category, confidence score and care recommendation. Confidence values must be copied directly from the live application; they must not be added manually.')
caption('Fig. 6. Sample garment image predictions with predicted fabric class and confidence. Insert actual LaundraAI screenshots here.')
para('The care recommendation module converts the predicted fabric category into transparent washing, drying, ironing and eco-care guidance. This recommendation is retrieved from the structured care knowledge base rather than being treated as a direct CNN output. Consequently, the system can explain why a gentle, cold or low-heat instruction is suggested.')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid';
for c,h in zip(t.rows[0].cells,['Fabric','Washing','Drying','Ironing','Eco tip']):c.text=h
for r in [['Cotton','30°C normal/gentle','Air dry','Medium heat','Full loads; low temperature'],['Polyester','30°C gentle','Air dry or low tumble','Low heat','Air dry where practical'],['Denim','Cold/30°C inside out','Air dry','Medium heat','Wash only when needed'],['Wool','Cold wool/hand wash','Dry flat','Low heat with cloth','Air between wears'],['Silk','Cold gentle/hand wash','Air dry away from sun','Low heat inside out','Short low-temperature care']]:
 for c,x in zip(t.add_row().cells,r):c.text=x
caption('Table III. Care recommendation results generated from the fabric knowledge base.')
para('User evaluation has not yet been conducted. After collecting actual Google Form responses, insert the satisfaction or usability chart as Fig. 7 and report the participant count, question scale and measured result. The current manuscript must not claim user-satisfaction findings before responses are collected.')
caption('Fig. 7. User satisfaction and usability evaluation. Insert Google Forms summary chart after data collection.')
para('Overall, the findings demonstrate that the proposed system can connect image-based fabric classification with explainable care guidance in one workflow. The confusion between some fabric types confirms that performance is dependent on image quality and dataset coverage. Future work should include additional independently collected garment images, blended-fabric labels, real-world lighting tests, care-label OCR and a formal user study.')
out=D/'LaundraAI_Results_and_Discussion_Manuscript_Style.docx';doc.save(out);print(out)

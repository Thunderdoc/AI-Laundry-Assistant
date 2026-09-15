import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'LaundraAI_Results_and_Discussion.docx'
metrics=json.loads((ROOT/'backend'/'models'/'fabric_mobilenetv2.manifest.json').read_text())
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos')
for n,size in [('Title',18),('Heading 1',14),('Heading 2',12)]:
 s=styles[n]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
def shade(cell,color):
 tc=cell._tc; pr=tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),color); pr.append(sh)
def set_cell(cell,text,bold=False,color=None):
 cell.text=''; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(9); 
 if color:r.font.color.rgb=RGBColor(*color)
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers,rows):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
 for c,h in zip(t.rows[0].cells,headers): set_cell(c,h,True,(255,255,255));shade(c,'0C6247')
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for c,v in zip(cells,row): set_cell(c,v)
  if i%2: [shade(c,'EEF6F1') for c in cells]
 doc.add_paragraph().paragraph_format.space_after=Pt(4);return t
def caption(text):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(text);r.italic=True;r.font.size=Pt(9)
def heading(text,level=1): doc.add_heading(text,level)

p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('V RESULTS AND DISCUSSION')
p=doc.add_paragraph('This section reports evaluation of the deployed LaundraAI MobileNetV2 model on a held-out test set. Fabric-care guidance is generated separately by the documented recommendation knowledge base.');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
heading('A Fabric Classification Performance',1);doc.add_paragraph('Table I presents overall classification results from the exported model evaluation artifact.')
macro_p=sum(v['precision'] for v in metrics['per_class_metrics'].values())/len(metrics['per_class_metrics']);macro_r=sum(v['recall'] for v in metrics['per_class_metrics'].values())/len(metrics['per_class_metrics'])
table(['Metric','Result'],[['Test images',metrics['test_images']],['Accuracy',f"{metrics['test_accuracy']*100:.2f}%"],['Macro Precision',f'{macro_p*100:.2f}%'],['Macro Recall',f'{macro_r*100:.2f}%'],['Macro F1 score',f"{metrics['macro_f1']*100:.2f}%"],['Recorded inference time',f"{metrics['test_inference_seconds']:.4f} s"],['Model','MobileNetV2 transfer learning'],['Input size','224 × 224 pixels']]);caption('Table I Overall Classification Performance')
doc.add_paragraph('Fig. 5 Confusion Matrix. Insert the confusion-matrix screenshot exported from the LaundraAI Model Performance page here. The raw matrix is given in Table II.')
cm=metrics['confusion_matrix']; table(['Actual / Predicted']+cm['classes'],[[cls]+row for cls,row in zip(cm['classes'],cm['matrix'])]);caption('Table II Confusion Matrix Values')
heading('B Fabric wise Classification Results',1)
table(['Fabric class','Support','Precision','Recall','F1 score'],[[k.title(),v['support'],f"{v['precision']*100:.2f}%",f"{v['recall']*100:.2f}%",f"{v['f1_score']*100:.2f}%"] for k,v in metrics['per_class_metrics'].items()]);caption('Table III Fabric wise Classification Results')
heading('C Sample Prediction Results',1);doc.add_paragraph('Fig. 6 Actual garment images with predicted fabric and confidence. Insert three to five screenshots from actual LaundraAI analyses. Do not create or enter sample confidence values until the images are tested.')
table(['Sample ID','Actual garment image','Predicted fabric','Confidence','Verified label','Comment'],[['S1','Insert result screenshot','Fill from app','Fill from app','Verified label','—'],['S2','Insert result screenshot','Fill from app','Fill from app','Verified label','—'],['S3','Insert result screenshot','Fill from app','Fill from app','Verified label','—']]);caption('Table IV Sample Prediction Results Template')
heading('D Care Recommendation Results',1)
care=[['Cotton','30°C normal/gentle; mild detergent','Air dry where practical','Medium heat; steam optional','Use full loads and lowest suitable temperature.'],['Polyester','30°C gentle; mild liquid detergent','Air dry or tumble low','Low heat; pressing cloth','Prefer air drying when practical.'],['Denim','Cold/30°C gentle; wash inside out','Air dry; avoid prolonged heat','Medium heat','Wash only when needed to reduce fading.'],['Wool','Cold wool/hand wash; wool detergent','Dry flat; reshape damp','Low heat with cloth','Air between wears before washing.'],['Silk','Cold gentle/hand wash; silk detergent','Air dry away from sun','Low heat, inside out','Use short low-temperature care.']]
table(['Fabric','Washing','Drying','Ironing','Eco tip'],care);caption('Table V Care Recommendation Results')
heading('E User Evaluation',1);doc.add_paragraph('Google Form responses have not yet been collected. Fig. 7 must be added only after actual participants complete the form. Recommended questions cover upload ease, result clarity, recommendation clarity, visual design, confidence transparency and overall satisfaction using a five-point Likert scale.')
doc.add_paragraph('Fig. 7 User satisfaction and usability chart. Insert the exported Google Forms summary chart here after data collection.')
heading('Discussion',1)
doc.add_paragraph(f"The model achieved {metrics['test_accuracy']*100:.2f}% accuracy and {metrics['macro_f1']*100:.2f}% macro F1-score on the {metrics['test_images']}-image held-out test set. Denim achieved the strongest F1-score, whereas silk was the most challenging class. The error pattern is consistent with visually similar fabric surfaces being difficult to distinguish from ordinary photographs.")
doc.add_paragraph('The non-fabric class provides an additional safeguard against unrelated uploads. The CNN provides fabric probabilities; washing, drying and ironing recommendations are then retrieved from the structured care knowledge base. Therefore, the recommendations are transparent rule-based guidance rather than direct neural-network outputs. The prototype should be treated as planning and decision-support assistance, and manufacturer care labels remain the authoritative source.')
doc.save(OUT);print(OUT)
